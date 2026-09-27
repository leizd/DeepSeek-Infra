"""Text-upload parity probe.

Calls the unmodified ``extract_uploaded_file``. The RAG sqlite index write is
patched to a no-op so the probe does not touch the host ``.local-rag`` store;
the file-cache directory is a temporary root. Chunk vectors are dropped before
comparison because the oracle embedding pipeline is hash-or-ONNX.

Usage::

    python tasks/native-runtime/file_text_parity_probe.py \\
        --rust-example rust/target/debug/examples/file_text_parity_probe.exe
"""

from __future__ import annotations

import argparse
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

CASES: list[tuple[str, str, str, bytes]] = [
    ("hello", "notes.txt", "text/plain", b"hello\nworld"),
    ("html", "page.html", "text/html", b"<p>Hello</p><script>secret()</script>"),
    ("cjk", "页.txt", "text/plain", "第一页\n第二页".encode()),
    ("empty", "empty.txt", "text/plain", b""),
    ("blank", "blank.txt", "text/plain", b" \n\t "),
    ("binary", "blob.bin", "application/octet-stream", b"a\0b"),
    ("image", "scan.png", "image/png", b"\x89PNG"),
]

CHUNK_KEYS = ("index", "start", "end", "lineStart", "lineEnd", "text")


def _outcome(cache_dir: Path, call: Any) -> dict[str, Any]:
    from deepseek_infra.core.errors import AppError

    try:
        value = call()
    except AppError as exc:
        return {"error": {"message": str(exc), "code": exc.code.value, "status": exc.status}}
    file_id = str(value.get("fileId") or "")
    raw = json.loads((cache_dir / f"{file_id}.json").read_text(encoding="utf-8"))
    chunks = [
        {key: chunk.get(key) for key in CHUNK_KEYS}
        for chunk in raw.get("chunks") or []
        if isinstance(chunk, dict)
    ]
    return {"ok": value, "chunks": chunks}


DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:r><w:t>Hello</w:t></w:r>
      <w:r><w:tab/></w:r>
      <w:r><w:t>world</w:t></w:r>
    </w:p>
    <w:p><w:r><w:t>Line</w:t><w:br/><w:t>two</w:t></w:r></w:p>
    <w:tbl>
      <w:tr>
        <w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p><w:p><w:r><w:t>C</w:t></w:r></w:p></w:tc>
      </w:tr>
    </w:tbl>
  </w:body>
</w:document>
"""

HEADER_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:p><w:r><w:t>Header</w:t></w:r></w:p>
</w:hdr>
"""

SLIDE_HELLO = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree>
    <p:sp><p:txBody><a:p><a:r><a:t>  Hello  </a:t></a:r></a:p></p:txBody></p:sp>
    <p:sp><p:txBody><a:p><a:r><a:t>World</a:t></a:r></a:p></p:txBody></p:sp>
  </p:spTree></p:cSld>
</p:sld>
"""

SLIDE_EMPTY = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree></p:spTree></p:cSld>
</p:sld>
"""

SLIDE_LATER = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree>
    <p:sp><p:txBody><a:p><a:r><a:t>Later</a:t></a:r></a:p></p:txBody></p:sp>
  </p:spTree></p:cSld>
</p:sld>
"""

EMPTY_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:sectPr/></w:body>
</w:document>
"""


def _zip_bytes(entries: list[tuple[str, str]]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, text in entries:
            archive.writestr(name, text.encode("utf-8"))
    return buffer.getvalue()


def write_docx_fixtures(directory: Path) -> None:
    (directory / "sample.docx").write_bytes(
        _zip_bytes(
            [
                ("word/document.xml", DOCUMENT_XML),
                ("word/header1.xml", HEADER_XML),
                ("word/footer2.xml", HEADER_XML.replace("Header", "Footer")),
            ]
        )
    )
    (directory / "empty.docx").write_bytes(_zip_bytes([("word/document.xml", EMPTY_XML)]))
    (directory / "bad.docx").write_bytes(b"this is not a zip")
    (directory / "broken.docx").write_bytes(_zip_bytes([("word/document.xml", "<w:document>")]))
    (directory / "sample.pptx").write_bytes(
        _zip_bytes(
            [
                ("ppt/slides/slide10.xml", SLIDE_LATER),
                ("ppt/slides/slide2.xml", SLIDE_EMPTY),
                ("ppt/slides/slide1.xml", SLIDE_HELLO),
            ]
        )
    )
    (directory / "empty.pptx").write_bytes(_zip_bytes([("ppt/slides/slide1.xml", SLIDE_EMPTY)]))
    (directory / "bad.pptx").write_bytes(b"this is not a zip")
    (directory / "broken.pptx").write_bytes(_zip_bytes([("ppt/slides/slide1.xml", "<p:sld>")]))
    write_xlsx_fixtures(directory)
    write_pdf_fixtures(directory)
    write_epub_fixtures(directory)


XLSX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""

XLSX_WORKBOOK = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="Strings" sheetId="1" r:id="rId1"/>
    <sheet name="Notes" sheetId="2" r:id="rId2"/>
  </sheets>
</workbook>
"""

XLSX_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
</Relationships>
"""

XLSX_SHARED = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="4" uniqueCount="4">
  <si><t>Hello</t></si>
  <si><r><t>Rich</t></r><r><t xml:space="preserve"> tail</t></r></si>
  <si><t>prex005F_post</t></si>
  <si><t>Base</t><rPh sb="0" eb="4"><t>NO</t></rPh></si>
</sst>
"""

XLSX_STRING_SHEET = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1">
      <c r="A1" t="s"><v>0</v></c>
      <c r="B1" t="s"><v>1</v></c>
      <c r="C1" t="s"><v>2</v></c>
      <c r="D1" t="s"><v>3</v></c>
      <c r="E1" t="n"><v>2</v></c>
    </row>
  </sheetData>
</worksheet>
"""

XLSX_NOTES_SHEET = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1"><c r="A1" t="inlineStr"><is><t>Tail</t></is></c></row>
  </sheetData>
</worksheet>
"""


def write_xlsx_fixtures(directory: Path) -> None:
    import openpyxl
    from datetime import date, datetime

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("openpyxl returned no active sheet")
    sheet.title = "Data"
    sheet["A1"] = "Hello"
    sheet["B1"] = 2
    sheet["C1"] = "  Hi  "
    sheet["B3"] = 2.5
    sheet["C3"] = True
    sheet["D3"] = False
    sheet["E3"] = datetime(2020, 1, 2, 3, 4, 5)
    sheet["F3"] = date(2020, 6, 7)
    workbook.create_sheet("Blank")
    later = workbook.create_sheet("Later")
    later["A4"] = "Later"
    workbook.save(directory / "sample.xlsx")

    empty = openpyxl.Workbook()
    empty.save(directory / "empty.xlsx")
    (directory / "bad.xlsx").write_bytes(b"this is not a zip")
    (directory / "broken.xlsx").write_bytes(_zip_bytes([("[Content_Types].xml", "<Types>")]))
    (directory / "nocontent.xlsx").write_bytes(
        _zip_bytes([("xl/worksheets/sheet1.xml", XLSX_NOTES_SHEET)])
    )
    (directory / "shared.xlsx").write_bytes(
        _zip_bytes(
            [
                ("[Content_Types].xml", XLSX_CONTENT_TYPES),
                ("xl/workbook.xml", XLSX_WORKBOOK),
                ("xl/_rels/workbook.xml.rels", XLSX_RELS),
                ("xl/sharedStrings.xml", XLSX_SHARED),
                ("xl/worksheets/sheet1.xml", XLSX_STRING_SHEET),
                ("xl/worksheets/sheet2.xml", XLSX_NOTES_SHEET),
            ]
        )
    )


def _pdf_bytes(pages: list[list[str]]) -> bytes:
    bodies: dict[int, bytes] = {}
    next_id = 1

    def alloc(body: bytes) -> int:
        nonlocal next_id
        number = next_id
        next_id += 1
        bodies[number] = body
        return number

    font = alloc(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pages_id = alloc(b"")
    catalog = alloc(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("ascii"))
    page_ids = []
    for lines in pages:
        commands = ["BT", "/F1 12 Tf", "72 720 Td"]
        for index, line in enumerate(lines):
            if index:
                commands.append("0 -16 Td")
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj")
        commands.append("ET")
        stream = ("\n".join(commands) + "\n").encode("latin1")
        content = alloc(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream")
        page = alloc(
            (
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
                f"/Contents {content} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>"
            ).encode("ascii")
        )
        page_ids.append(page)
    kids = " ".join(f"{page} 0 R" for page in page_ids)
    bodies[pages_id] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("ascii")
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in range(1, next_id):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode("ascii")
        body = bodies[number]
        out += body
        if not body.endswith(b"\n"):
            out += b"\n"
        out += b"endobj\n"
    xref = len(out)
    out += f"xref\n0 {next_id}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for number in range(1, next_id):
        out += f"{offsets[number]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {next_id} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def write_epub_fixtures(directory: Path) -> None:
    (directory / "sample.epub").write_bytes(
        _zip_bytes(
            [
                ("b.xhtml", "<h1>Later</h1>"),
                ("nav.xhtml", "<p>SECRET NAV</p>"),
                ("sub/toc.xhtml", "<p>SECRET TOC</p>"),
                ("a.xhtml", "<p>Hello epub</p><script>secret()</script>"),
                ("C.HTML", "<p>Upper</p>"),
            ]
        )
    )
    (directory / "empty.epub").write_bytes(_zip_bytes([("nav.xhtml", "<p>SECRET NAV</p>")]))
    (directory / "bad.epub").write_bytes(b"this is not a zip")


def write_pdf_fixtures(directory: Path) -> None:
    import fitz

    (directory / "sample.pdf").write_bytes(
        _pdf_bytes([["Hello pdf", "World", "A (b) c"], ["Second"]])
    )
    (directory / "blank.pdf").write_bytes(_pdf_bytes([["   "]]))
    (directory / "bad.pdf").write_bytes(b"this is not a pdf")
    (directory / "header.pdf").write_bytes(b"%PDF-1.4")
    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 72), "Hello pdf")
    first.insert_text((72, 100), "Second line")
    second = document.new_page()
    second.insert_text((72, 72), "Page two")
    (directory / "drawn.pdf").write_bytes(document.tobytes())
    document.close()


def oracle_report(fixtures: Path) -> dict[str, Any]:
    from deepseek_infra.infra.rag import files as rag_files
    from deepseek_infra.infra.rag import local_rag

    saved_dir = rag_files.FILE_CACHE_DIR
    saved_index = local_rag.index_file_payload
    root = Path(tempfile.mkdtemp(prefix="file-text-probe-"))
    cache_dir = root / ".file-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    rag_files.FILE_CACHE_DIR = cache_dir
    local_rag.index_file_payload = lambda *args, **kwargs: 0
    try:
        report: dict[str, Any] = {}
        for label, filename, content_type, data in CASES:
            report[label] = _outcome(
                cache_dir,
                lambda filename=filename, content_type=content_type, data=data: rag_files.extract_uploaded_file(
                    filename, content_type, data
                ),
            )
        docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        for label, filename in (
            ("docx", "sample.docx"),
            ("docx-empty", "empty.docx"),
            ("docx-bad", "bad.docx"),
            ("docx-xml", "broken.docx"),
        ):
            data = (fixtures / filename).read_bytes()
            report[label] = _outcome(
                cache_dir,
                lambda filename=filename, data=data: rag_files.extract_uploaded_file(
                    filename, docx_type, data
                ),
            )
        pptx_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        for label, filename in (
            ("pptx", "sample.pptx"),
            ("pptx-empty", "empty.pptx"),
            ("pptx-bad", "bad.pptx"),
            ("pptx-xml", "broken.pptx"),
        ):
            data = (fixtures / filename).read_bytes()
            report[label] = _outcome(
                cache_dir,
                lambda filename=filename, data=data: rag_files.extract_uploaded_file(
                    filename, pptx_type, data
                ),
            )
        xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        for label, filename in (
            ("xlsx", "sample.xlsx"),
            ("xlsx-empty", "empty.xlsx"),
            ("xlsx-bad", "bad.xlsx"),
            ("xlsx-xml", "broken.xlsx"),
            ("xlsx-nocontent", "nocontent.xlsx"),
            ("xlsx-shared", "shared.xlsx"),
        ):
            data = (fixtures / filename).read_bytes()
            report[label] = _outcome(
                cache_dir,
                lambda filename=filename, data=data: rag_files.extract_uploaded_file(
                    filename, xlsx_type, data
                ),
            )
        pdf_type = "application/pdf"
        epub_type = "application/epub+zip"
        for label, filename in (
            ("pdf", "sample.pdf"),
            ("pdf-blank", "blank.pdf"),
            ("pdf-bad", "bad.pdf"),
            ("pdf-header", "header.pdf"),
            ("pdf-drawn", "drawn.pdf"),
            ("epub", "sample.epub"),
            ("epub-empty", "empty.epub"),
            ("epub-bad", "bad.epub"),
            ("ocr-image", "ocr-hello.png"),
            ("ocr-pdf", "ocr-scan.pdf"),
        ):
            data = (fixtures / filename).read_bytes()
            if filename.endswith(".epub"):
                content_type = epub_type
            elif filename.endswith(".png"):
                content_type = "image/png"
            else:
                content_type = pdf_type
            ocr_enabled = filename.startswith("ocr-")
            report[label] = _outcome(
                cache_dir,
                lambda filename=filename, content_type=content_type, data=data, ocr_enabled=ocr_enabled: rag_files.extract_uploaded_file(
                    filename, content_type, data, ocr_enabled=ocr_enabled
                ),
            )
        return report
    finally:
        rag_files.FILE_CACHE_DIR = saved_dir
        local_rag.index_file_payload = saved_index
        shutil.rmtree(root, ignore_errors=True)


def _hello_ocr_png() -> bytes:
    import io

    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 96)
    image = Image.new("RGB", (360, 140), "white")
    draw = ImageDraw.Draw(image)
    draw.text((8, 12), "HELLO", fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def write_ocr_fixtures(directory: Path) -> None:
    import io

    from PIL import Image

    png = _hello_ocr_png()
    (directory / "ocr-hello.png").write_bytes(png)
    pdf = io.BytesIO()
    Image.open(io.BytesIO(png)).convert("RGB").save(pdf, format="PDF")
    (directory / "ocr-scan.pdf").write_bytes(pdf.getvalue())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--rust-json", type=Path)
    source.add_argument("--rust-example", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    # pix2tex's text changes between processes, so a byte compare is not a
    # property of the image. Pin the formula engine to a command that fails
    # and let Tesseract / Windows OCR produce the stable text.
    os.environ["OCR_FORMULA_CMD"] = "cmd /c exit 1"
    fixtures = Path(tempfile.mkdtemp(prefix="file-text-docx-"))
    try:
        write_docx_fixtures(fixtures)
        write_ocr_fixtures(fixtures)
        if args.rust_example is not None:
            environment = os.environ.copy()
            environment["FILE_TEXT_DOCX_DIR"] = str(fixtures)
            completed = subprocess.run(
                [str(args.rust_example)],
                capture_output=True,
                text=True,
                encoding="utf-8",
                check=False,
                env=environment,
            )
            if completed.returncode != 0:
                print(completed.stderr, file=sys.stderr)
                return 2
            native = json.loads(completed.stdout)
        else:
            native = json.loads(Path(args.rust_json).read_text(encoding="utf-8"))
        oracle = oracle_report(fixtures)
    finally:
        shutil.rmtree(fixtures, ignore_errors=True)
    problems: list[str] = []
    compared = 0
    for key in sorted(set(oracle) | set(native)):
        compared += 1
        if oracle.get(key) != native.get(key):
            problems.append(f"{key}: oracle={oracle.get(key)!r} native={native.get(key)!r}")
    report = {"cases_compared": compared, "problems": problems, "result": "PASS" if not problems else "FAIL"}
    if args.report is not None:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))
    return 0 if not problems else 1


if __name__ == "__main__":
    raise SystemExit(main())
